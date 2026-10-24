"""Defensive, corpus-relative document ingestion. Never executes source files."""
import csv
import re
from pathlib import Path

TEXT_TYPES={'.txt','.log','.py','.md','.json','.yaml','.yml'}
IMAGE_TYPES={'.png','.jpg','.jpeg','.tif','.tiff'}


def _read_file(path,ocr):
    ext=path.suffix.lower()
    if ext in TEXT_TYPES:return path.read_text(encoding='utf8',errors='replace')
    if ext=='.csv':
        with path.open(encoding='utf-8-sig',errors='replace',newline='') as f:
            rows=csv.DictReader(f)
            return '\n'.join(' | '.join(f'{k}: {v}' for k,v in row.items() if v is not None) for row in rows)
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:raise ValueError('Encrypted PDF; no password supplied')
        return '\n'.join(f'[page {i+1}]\n{page.extract_text() or ""}' for i,page in enumerate(reader.pages))
    if ext=='.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        doc=Document(path);parts=[]
        for element in doc.element.body.iterchildren():
            if element.tag.endswith('}p'):parts.append(Paragraph(element,doc).text)
            elif element.tag.endswith('}tbl'):
                table=Table(element,doc)
                parts.extend(' | '.join(cell.text for cell in row.cells) for row in table.rows)
        return '\n'.join(parts)
    if ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False,keep_links=False);parts=[]
        try:
            for sheet in book.worksheets:
                rows=sheet.iter_rows(values_only=True);headers=next(rows,())
                parts.append('[sheet '+sheet.title+'] '+ ' | '.join(str(x or '') for x in headers))
                for row in rows:
                    parts.append(' | '.join(f'{headers[i] if i<len(headers) else i}: {value}' for i,value in enumerate(row) if value is not None))
        finally:book.close()
        return '\n'.join(parts)
    if ext in IMAGE_TYPES:
        if ocr is None:raise ValueError('Vision OCR callback not configured')
        return ocr(path)
    return None


def _withdrawn(path,text):
    if re.search(r'(?:^|[_ .-])(WITHDRAWN|SUPERSEDED|OBSOLETE|RETIRED)(?:[_ .-]|$)',path.stem,re.I):return True
    return bool(re.search(r'(?im)^\s*(?:document\s+)?status\s*[:=]\s*(?:withdrawn|superseded|obsolete|retired)\b',text[:3000]))


def read_corpus(root,ocr=None):
    root=Path(root).resolve();documents=[];errors=[]
    for path in sorted(root.rglob('*')):
        rel=path.relative_to(root).as_posix()
        try:
            if path.is_symlink():raise ValueError('Symlink excluded from corpus')
            if not path.is_file():continue
            if not path.stat().st_mode & 0o444:raise PermissionError('File has no read permission bits')
            text=_read_file(path,ocr)
            if text is None:continue
            if _withdrawn(path,text):
                errors.append({'file':rel,'reason':'withdrawn'});continue
            if text.strip():documents.append({'file':rel,'text':text,'kind':path.suffix.lower()})
        except Exception as exc:errors.append({'file':rel,'reason':str(exc)})
    return documents,errors

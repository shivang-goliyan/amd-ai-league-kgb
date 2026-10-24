import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
APP=Path(__file__).resolve().parents[1]/'app';sys.path.insert(0,str(APP))
import documents

class FormatTests(unittest.TestCase):
    def test_pdf_text_and_encryption(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);w=PdfWriter();page=w.add_blank_page(width=300,height=300)
            font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
            stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 260 Td (Temperature limit: 87 C) Tj ET');page[NameObject('/Contents')]=w._add_object(stream)
            with (root/'normal.pdf').open('wb') as f:w.write(f)
            w.encrypt('not-provided-to-reader')
            with (root/'locked.pdf').open('wb') as f:w.write(f)
            docs,errors=documents.read_corpus(root)
            self.assertEqual([x['file'] for x in docs],['normal.pdf']);self.assertIn('87',docs[0]['text']);self.assertTrue(any('Encrypted' in x['reason'] for x in errors))

    def test_docx_reads_table_between_paragraphs(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);doc=Document();doc.add_paragraph('Roadmap start');t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='Product';t.cell(0,1).text='Sampling';t.cell(1,0).text='AX-71';t.cell(1,1).text='Q2 FY28';doc.add_paragraph('Roadmap end');doc.save(root/'roadmap.docx')
            docs,_=documents.read_corpus(root);text=docs[0]['text'];self.assertIn('AX-71',text);self.assertIn('Q2 FY28',text);self.assertLess(text.index('Roadmap start'),text.index('Q2 FY28'));self.assertLess(text.index('Q2 FY28'),text.index('Roadmap end'))

    def test_xlsx_reads_every_sheet(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);book=Workbook();a=book.active;a.title='Parts';a.append(['Part','Lead time']);a.append(['FAN-672',12]);b=book.create_sheet('Revisions');b.append(['Board','Revision']);b.append(['AX-71','REV-F3']);book.save(root/'parts.xlsx');book.close()
            docs,_=documents.read_corpus(root);text=docs[0]['text']
            for value in ['Parts','Lead time: 12','Revisions','REV-F3']:self.assertIn(value,text)

if __name__=='__main__':unittest.main()
